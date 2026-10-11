#!/usr/bin/env python3
"""A folder like a real one: tools saved as Rev 1..3 and v1..v3_final, browser download copies "(1)",
"Copy of", the same file in two folders, dated saves, near-identical weekly dashboards, plus pages that
look alike but are different things (two cabins, two cities) and must stay separate. Also songs, videos
and pictures that are the same thing in different files."""
import os, random, shutil, subprocess, sys, json
from PIL import Image, ImageDraw, ImageFilter

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'Workbench')
shutil.rmtree(ROOT, ignore_errors=True)
R = random.Random(1998)

def P(*parts):
    p = os.path.join(ROOT, *parts)
    os.makedirs(os.path.dirname(p), exist_ok=True)
    return p

WORDS = {
    'cayenne': 'cayenne diesel coolant hatch strut bumper exhaust tdi turbo injector glow plug filter oil transmission mileage service interval brake rotor caliper pad sensor adblue urea tank pump gasket seal hose clamp radiator thermostat battery alternator starter tire tread alignment suspension air spring compressor valve module fuse relay harness connector torque socket wrench jack stand lift dealer invoice receipt warranty recall bulletin parts catalog supplier shipping'.split(),
    'poker': 'tournament stack blinds ante bounty satellite downgrade campaign bankroll buyin payout ladder bubble icm chip leader shove fold call raise flop turn river range equity outs pot odds implied position button cutoff hijack lojack blind defend steal squeeze limp isolate continuation barrel bluff value thin river check jam pushfold table final schedule flight structure level break registration'.split(),
    'crucible': 'agent debate critic judge verdict adversarial decision engine scenario option tradeoff argument rebuttal evidence weight rubric score consensus dissent round prompt persona skeptic optimist auditor transcript summary confidence assumption risk mitigation framework ledger memo recommendation alternative constraint objective stakeholder premise'.split(),
    'cruise': 'coral duchess deck cabin balcony starboard port midship stateroom steward dining venue buffet specialty reservation patter excursion tender embarkation disembarkation muster drill lido atrium piazza theater spa sanctuary pool sunrise sunset itinerary island beach snorkel shore coral cay cozumel ocean fare premier plus standard package'.split(),
    'fantasy': 'quarterback receiver running tight defense kicker matchup projection points bench flex waiver trade roster lineup starter injury questionable doubtful touchdown yards reception target snap share red zone opponent standings playoff seed bye'.split(),
    'moe': 'expert router gating token throughput memory bandwidth channel dimm socket server rack quantization layer offload batch context prompt prefill decode latency vram ram ddr4 ddr5 epyc xeon motherboard power supply cooling fan noise benchmark llama mixtral deepseek qwen inference cluster'.split(),
    'italy': 'rome venice florence gondola piazza basilica museum gelato espresso trattoria vaporetto canal bridge palazzo duomo uffizi colosseum forum pantheon vatican train station hotel concierge reservation dinner tasting wine harvest tour guide ticket'.split(),
    'common': 'plan notes budget total cost estimate option choice detail summary schedule time date day morning afternoon evening update review status open closed done pending next step priority list item check'.split(),
}
FILLER = 'the a of and to in for on with at by from is are this that it as be was were or an'.split()

def para(topic, n, rnd, extra=()):
    vocab = WORDS[topic] + WORDS['common'] + list(extra)
    out = []
    for _ in range(n):
        w = rnd.choice(vocab) if rnd.random() < 0.7 else rnd.choice(FILLER)
        out.append(w)
    s = ' '.join(out)
    return s[0].upper() + s[1:] + '.'

def page(title, body_html, script=''):
    return f'''<!DOCTYPE html><html lang="en"><head><meta charset="utf-8"><title>{title}</title>
<style>body{{font-family:Manrope,Arial,sans-serif;background:#fbf6ec;color:#222}}h1{{font-family:Fraunces,serif;color:#a0522d}}.card{{border:1px solid #d9c9a3;padding:12px;margin:10px 0}}</style></head>
<body><h1>{title}</h1>{body_html}<script>{script}</script></body></html>'''

def write(path, text):
    with open(path, 'w', encoding='utf-8') as f: f.write(text)

def base_words(topic, n, seed, extra=()):
    rnd = random.Random(seed)
    vocab = WORDS[topic] + WORDS['common'] + list(extra)
    return [rnd.choice(vocab) for _ in range(n)]

def revision_set(topic, seed, revs, title_fn, name_fn, folder, grow=18, extra=()):
    """revs of one tool: each keeps most of the last one's text and adds a little"""
    rnd = random.Random(seed)
    paras = [para(topic, 40, rnd, extra) for _ in range(8)]
    out = []
    for i in range(1, revs + 1):
        if i > 1:
            paras = paras[:]  # keep almost everything
            paras[rnd.randrange(len(paras))] = para(topic, 40, rnd, extra)   # one paragraph rewritten
            for _ in range(2): paras.append(para(topic, grow, rnd, extra))   # new features
        body = ''.join(f'<div class="card"><h2>Section {k + 1}</h2><p>{p}</p></div>' for k, p in enumerate(paras))
        js = "const ui = document.querySelector('h1'); ui.style.cssText = 'color:#a0522d'; const tpl = `<div style=\"padding:4px\">${'{'}name{'}'}</div>`; localStorage.setItem('rev', '" + str(i) + "');"
        path = P(*folder, name_fn(i))
        write(path, page(title_fn(i), body, js))
        out.append(path)
        os.utime(path, (1_790_000_000 + i * 86400, 1_790_000_000 + i * 86400))
    return out

# ---- tools saved as revisions ----
revision_set('cayenne', 11, 3, lambda i: f'Cayenne 958 Tracker — Rev {i}', lambda i: f'Cayenne_958_Tracker_Rev{i}.html' if i < 3 else 'Cayenne 958 Tracker — Rev 3.html', ('Tools', 'Cayenne'))
revision_set('poker', 12, 3, lambda i: f'Deep Run Tournament Trainer v{i}', lambda i: f'Deep Run Tournament Trainer v{i}.html' if i < 3 else 'Deep_Run_Tournament_Trainer_v3_final.html', ('Tools', 'Poker'))
revision_set('crucible', 13, 3, lambda i: f'Crucible · Adversarial Decision Engine · Rev {i + 3}', lambda i: f'Crucible Rev {i + 3}.html' if i < 3 else 'Crucible_Rev6.html', ('Tools', 'Crucible'))
revision_set('moe', 14, 2, lambda i: 'MoE 256GB Build Study', lambda i: f'MoE_256GB_Build_Study_{i + 2}.html', ('Downloads',))

# ---- distinct tools on the same topics (these must stay separate, and make the rings) ----
def tool(topic, seed, title, *path, extra=()):
    rnd = random.Random(seed)
    body = ''.join(f'<div class="card"><p>{para(topic, 45, rnd, extra)}</p></div>' for _ in range(6))
    write(P(*path), page(title, body))
tool('cayenne', 21, 'Cayenne Auction Board', 'Tools', 'Cayenne', 'Cayenne Auction Board.html')
tool('cayenne', 22, 'Cayenne Rear-End Project Workspace', 'Tools', 'Cayenne', 'Cayenne Rear-End Project Workspace.html')
tool('poker', 23, 'The Corner — Tournament Tracker', 'Tools', 'Poker', 'The Corner.html')
tool('poker', 24, 'WSOP Downgrade Calculator', 'Tools', 'Poker', 'WSOP Downgrade Calculator.html')
tool('poker', 25, 'MSPT Campaign Planner', 'Tools', 'Poker', 'MSPT Campaign Planner.html')
tool('crucible', 26, 'Crucible Prompt Library', 'Tools', 'Crucible', 'Crucible Prompt Library.html')
tool('moe', 27, 'P40 Rig Setup Notes', 'Downloads', 'P40 Rig Setup Notes.html')

# cruise pages share a theme (and a lot of words) but are different tools
CRUISE_X = ['departure', '4:00PM', '5th', 'deck', '00PM']
for seed, title, name in [(31, 'Coral Duchess Meal Planner', 'Coral Duchess Meal Planner.html'), (32, 'Coral Duchess Packing & Boarding', 'Coral Duchess Packing & Boarding.html'),
                          (33, 'Coral Duchess Dining Schedule', 'Coral Duchess Dining Schedule.html'), (34, 'Coral Cay Picnic Plan', 'Coral Cay Picnic Plan.html')]:
    rnd = random.Random(seed)
    body = '<p>All aboard by 4:00PM. Muster on the 5th deck.</p>' + ''.join(f'<div class="card"><p>{para("cruise", 45, rnd)}</p></div>' for _ in range(6))
    write(P('Cruise', name), page(title, body))
# two cabin dossiers from one template: alike, but different cabins
template = [para('cruise', 40, random.Random(40)) for _ in range(4)]
for cabin, seed in [('B210', 41), ('B212', 42)]:
    rnd = random.Random(seed)
    own = [para('cruise', 40, rnd, extra=[cabin.lower(), 'aloha', 'forward', 'aft', 'seawalk', 'clearance', 'compass', 'heading', 'sun', 'shade', 'wind', 'noise', 'elevator', 'stairwell', 'laundry']) for _ in range(5)]
    body = f'<p>Cabin {cabin} on Aloha deck.</p>' + ''.join(f'<div class="card"><p>{p}</p></div>' for p in template + own)
    write(P('Cruise', f'Mini-Suite {cabin} Cabin Dossier.html'), page(f'Mini-Suite {cabin} Cabin Dossier', body))
# download duplicates of the meal planner: "(1)" next to it, and the same file in Downloads
shutil.copy(P('Cruise', 'Coral Duchess Meal Planner.html'), P('Cruise', 'Coral Duchess Meal Planner (1).html'))
shutil.copy(P('Cruise', 'Coral Duchess Meal Planner.html'), P('Downloads', 'Coral Duchess Meal Planner.html'))
shutil.copy(P('Cruise', 'Coral Duchess Dining Schedule.html'), P('Cruise', 'Copy of Coral Duchess Dining Schedule.html'))

# weekly dashboards: the same tool with this week's numbers
roster = 'mahomes kelce hill waddle pacheco mccaffrey jefferson lamb chase brown kittle tucker'.split()
for wk in (4, 5):
    rnd = random.Random(50 + wk)
    p1 = ' '.join(roster) + '. ' + ' '.join(WORDS['fantasy'])
    body = f'<p>{p1}</p>' + ''.join(f'<p>{para("fantasy", 12, rnd)}</p>' for _ in range(1))
    write(P('Fantasy', f'Fantasy Matchup Dashboard Week {wk}.html'), page(f'Fantasy Matchup Dashboard Week {wk}', body))
tool('fantasy', 55, 'Fantasy Draft Board', 'Fantasy', 'Fantasy Draft Board.html')

# an itinerary template used for two different cities: separate pages
it_template = [para('italy', 30, random.Random(60)) for _ in range(3)]
for city, seed in [('Rome', 61), ('Venice', 62)]:
    rnd = random.Random(seed)
    own = [para('italy', 45, rnd, extra=[city.lower()] * 8) for _ in range(5)]
    write(P('Italy', f'{city} Days.html'), page(f'{city} Days', ''.join(f'<p>{p}</p>' for p in it_template + own)))
# an edited copy: "Italy Trip - Copy" is the same trip with changes
rnd = random.Random(63)
trip = [para('italy', 40, rnd) for _ in range(6)]
write(P('Italy', 'Italy Trip.html'), page('Italy Harvest Trip', ''.join(f'<p>{p}</p>' for p in trip)))
trip2 = trip[:5] + [para('italy', 40, rnd)]
write(P('Italy', 'Italy Trip - Copy.html'), page('Italy Harvest Trip', ''.join(f'<p>{p}</p>' for p in trip2)))
os.utime(P('Italy', 'Italy Trip - Copy.html'), (1_790_500_000, 1_790_500_000))
os.utime(P('Italy', 'Italy Trip.html'), (1_790_000_000, 1_790_000_000))

# a revision saved under a slightly different name, somewhere else: "Cayenne Fuel Log" -> "Cayenne Diesel Fuel Log"
rnd = random.Random(90)
fuel = [para('cayenne', 40, rnd, extra=['fuel', 'gallons', 'mpg', 'fillup']) for _ in range(8)]
write(P('Tools', 'Cayenne', 'Cayenne Fuel Log.html'), page('Cayenne Fuel Log', ''.join(f'<p>{x}</p>' for x in fuel)))
fuel2 = fuel[:6] + [para('cayenne', 40, rnd, extra=['fuel', 'gallons', 'mpg', 'fillup']) for _ in range(2)]
write(P('Downloads', 'Cayenne Diesel Fuel Log.html'), page('Cayenne Diesel Fuel Log', ''.join(f'<p>{x}</p>' for x in fuel2)))
os.utime(P('Tools', 'Cayenne', 'Cayenne Fuel Log.html'), (1_790_000_000, 1_790_000_000))
os.utime(P('Downloads', 'Cayenne Diesel Fuel Log.html'), (1_790_300_000, 1_790_300_000))

# two different "notes.html" in two folders: generic name, different content -> separate
tool('cayenne', 70, 'Notes', 'Tools', 'Cayenne', 'notes.html')
tool('poker', 71, 'Notes', 'Tools', 'Poker', 'notes.html')

# a website folder and a backup copy of it: website pages never fold
for site in ('Site', 'Site copy'):
    write(P(site, 'index.html'), page('Homebrew AI Club', '<p>Welcome to the club. Meetings, builds and benchmarks.</p><a href="about.html">About</a>'))
    write(P(site, 'about.html'), page('About the Homebrew AI Club', '<p>We build local inference rigs and share notes.</p>'))

# ---- documents ----
from docx import Document
def docx(path, title, paras):
    d = Document(); d.core_properties.title = title; d.add_heading(title, 1)
    for p in paras: d.add_paragraph(p)
    d.save(path)
spec = [para('cayenne', 40, random.Random(80)) for _ in range(6)]
docx(P('Docs', 'Spec Review.docx'), 'Spec Review', spec)
docx(P('Docs', 'Spec Review v2.docx'), 'Spec Review', spec[:5] + [para('cayenne', 40, random.Random(81))])
from openpyxl import Workbook
wb = Workbook(); ws = wb.active; ws.title = 'Budget'
for r, row in enumerate([['Item', 'Cost'], ['Pizza', 12.99], ['Soda', 1.49], ['Cables', 4.5]], 1):
    for c, v in enumerate(row, 1): ws.cell(r, c, v)
wb.save(P('Docs', 'LAN Party Budget.xlsx'))
shutil.copy(P('Docs', 'LAN Party Budget.xlsx'), P('Docs', 'LAN Party Budget (1).xlsx'))

# ---- songs, videos, pictures ----
def ff(*args):
    subprocess.run(['ffmpeg', '-v', 'error', '-y', *args], check=True)
meta = ['-metadata', 'artist=The Hypertext Kids', '-metadata', 'title=Dial Tone Serenade', '-metadata', 'album=Under Construction']
ff('-f', 'lavfi', '-i', 'sine=frequency=440:sample_rate=44100:duration=6', *meta, '-c:a', 'libmp3lame', '-b:a', '128k', P('Music', 'The Hypertext Kids - Dial Tone Serenade.mp3'))
ff('-f', 'lavfi', '-i', 'sine=frequency=440:sample_rate=44100:duration=6', *meta, '-c:a', 'flac', P('Music', 'Lossless', 'Dial Tone Serenade.flac'))
shutil.copy(P('Music', 'The Hypertext Kids - Dial Tone Serenade.mp3'), P('Music', 'Backup', 'The Hypertext Kids - Dial Tone Serenade.mp3'))
ff('-f', 'lavfi', '-i', 'sine=frequency=660:sample_rate=44100:duration=5', '-metadata', 'artist=The Hypertext Kids', '-metadata', 'title=Packet Loss', '-metadata', 'album=Under Construction', '-c:a', 'libmp3lame', '-b:a', '128k', P('Music', 'The Hypertext Kids - Packet Loss.mp3'))
V = ['-f', 'lavfi', '-i', 'testsrc2=size=320x240:rate=24:duration=4']
ff(*V, '-c:v', 'libvpx-vp9', '-deadline', 'realtime', '-cpu-used', '8', '-b:v', '200k', P('Videos', 'Oil Change S01E01.webm'))
ff(*V, '-c:v', 'libvpx-vp9', '-deadline', 'realtime', '-cpu-used', '8', '-b:v', '300k', P('Videos', 'Old Rips', 'Oil Change S01E01.mkv'))
ff('-f', 'lavfi', '-i', 'testsrc=size=320x240:rate=24:duration=3', '-c:v', 'libvpx-vp9', '-deadline', 'realtime', '-cpu-used', '8', P('Videos', 'Oil Change S01E02.webm'))
img = Image.new('RGB', (1200, 800), (40, 120, 200))
d = ImageDraw.Draw(img)
for i in range(30):
    x, y = R.randrange(1200), R.randrange(800)
    d.ellipse([x, y, x + R.randrange(40, 200), y + R.randrange(40, 200)], fill=(R.randrange(256), R.randrange(256), R.randrange(256)))
img = img.filter(ImageFilter.GaussianBlur(2))
img.save(P('Pictures', 'beach.jpg'), quality=90)
shutil.copy(P('Pictures', 'beach.jpg'), P('Pictures', 'beach (1).jpg'))
img.resize((600, 400)).save(P('Pictures', 'beach_small.jpg'), quality=80)
img2 = Image.new('RGB', (1200, 800), (250, 120, 40))
d2 = ImageDraw.Draw(img2)
d2.rectangle([100, 500, 1100, 800], fill=(20, 20, 60))
img2.save(P('Pictures', 'sunset.jpg'), quality=90)

n = sum(len(f) for _, _, f in os.walk(ROOT))
print('made', n, 'files in', ROOT)
