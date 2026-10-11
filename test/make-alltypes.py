#!/usr/bin/env python3
"""Every kind of file LANscape opens, for test/viewers.js and test/spider2.js: pictures in three albums (one with
camera details), a website folder with its own logo, a slide deck, a ZIP, a font, a binary, a PDF spec sheet whose
key words are on its later pages, two WebM videos, songs, and Word, Excel, Markdown and text documents."""
import os, shutil, subprocess, zipfile, random
from PIL import Image, ImageDraw
from PIL.TiffImagePlugin import IFDRational

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'AllTypes')
shutil.rmtree(ROOT, ignore_errors=True)
R = random.Random(1999)


def P(*parts):
    p = os.path.join(ROOT, *parts)
    os.makedirs(os.path.dirname(p), exist_ok=True)
    return p


def write(path, text):
    with open(path, 'w', encoding='utf-8') as f:
        f.write(text)


# ---- pictures: three albums (Garage, Vacation 2019, Tutorials) ----
def pic(path, w, h, color, exif=None, fmt=None):
    im = Image.new('RGB', (w, h), color)
    d = ImageDraw.Draw(im)
    for i in range(14):
        x, y = R.randrange(w), R.randrange(h)
        r = R.randrange(w // 20, w // 6)
        d.ellipse([x, y, x + r, y + r], fill=(R.randrange(256), R.randrange(256), R.randrange(256)))
    kw = {}
    if (fmt or path.rsplit('.', 1)[-1].lower()) in ('jpg', 'jpeg'):
        kw['quality'] = 86
        if exif is not None:
            kw['exif'] = exif
    im.save(path, **kw)


ex = Image.Exif()
ex[271] = 'Kodak'
ex[272] = 'DC40'
sub = ex.get_ifd(0x8769)
sub[33434] = IFDRational(1, 250)      # 1/250 s
sub[33437] = IFDRational(56, 10)      # f/5.6
sub[36867] = '2019:07:04 14:32:10'    # taken
pic(P('Pictures', 'Vacation 2019', 'Grand Turk Beach.jpg'), 1600, 1067, (40, 140, 210), exif=ex)
pic(P('Pictures', 'Vacation 2019', 'Sunset Over Cozumel.jpg'), 1200, 800, (240, 120, 40))
pic(P('Pictures', 'Garage', 'Cayenne Front.jpg'), 1024, 768, (90, 90, 100))
pic(P('Tutorials', 'Wiring Diagram.png'), 800, 600, (250, 250, 240))

# ---- a website folder: its picture belongs to the website, not the Photo Album ----
logo = Image.new('RGB', (200, 80), (0, 0, 128))
ImageDraw.Draw(logo).rectangle([10, 10, 190, 70], outline=(255, 255, 0), width=4)
logo.save(P('Homepage', 'logo.png'))
write(P('Homepage', 'index.html'), '<!DOCTYPE html><html><head><title>Dale\'s Garage Homepage</title></head><body bgcolor="#ffffcc"><img src="logo.png" width="200" height="80" alt="logo"><h1>Welcome to my garage homepage</h1><p>Project cars, diesel tips and the occasional road trip.</p><p><a href="about.html">About me</a></p></body></html>')
write(P('Homepage', 'about.html'), '<!DOCTYPE html><html><head><title>About Dale</title></head><body><h1>About me</h1><p>I fix old cars and build web pages by hand.</p><p><a href="index.html">Home</a></p></body></html>')

# ---- pages, one of which builds its words with JavaScript (the Spider reads those) ----
write(P('Pages', 'Torque Specs.html'), '<!DOCTYPE html><html><head><title>Torque Specs</title></head><body><h1>Torque Specs</h1><div id="t"></div><script>document.getElementById("t").innerHTML = ["Lug nuts 118 ft-lb", "Caliper bolts 63 ft-lb", "Drain plug 37 ft-lb"].map(function (x) { return "<p>" + x + "</p>"; }).join("");</script></body></html>')
write(P('Pages', 'Road Trip Packing List.html'), '<!DOCTYPE html><html><head><title>Road Trip Packing List</title></head><body><h1>Packing list</h1><ul><li>Jumper cables</li><li>Tire gauge</li><li>Snacks</li></ul></body></html>')

# ---- a slide deck with three slides ----
from pptx import Presentation
from pptx.util import Inches, Pt
prs = Presentation()
s = prs.slides.add_slide(prs.slide_layouts[0])
s.shapes.title.text = 'Garage Tour'
s.placeholders[1].text = 'A walk around the shop, 1999'
for title, bullets in [('The Lift', ['Two-post, 9,000 lb', 'Installed last spring']), ('The Toolbox', ['Metric and SAE sockets', 'A torque wrench that clicks'])]:
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = title
    tf = s.placeholders[1].text_frame
    tf.text = bullets[0]
    for b in bullets[1:]:
        tf.add_paragraph().text = b
prs.save(P('Slides', 'Garage Tour.pptx'))

# ---- a ZIP with three files ----
with zipfile.ZipFile(P('Archives', 'project-backup.zip'), 'w', zipfile.ZIP_DEFLATED) as z:
    z.writestr('readme.txt', 'Project backup from the old 486. Restore these files to C:\\PROJECTS.\n')
    z.writestr('main.py', 'print("hello from 1999")\n')
    z.writestr('notes.md', '# Notes\n\n- remember to defrag\n')

# ---- a font and a binary ----
shutil.copy('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf', P('Fonts', 'DejaVuSans-Bold.ttf'))
with open(P('Misc', 'firmware.bin'), 'wb') as f:
    f.write(bytes(R.randrange(256) for _ in range(4096)))

# ---- a PDF spec sheet whose key words are on pages 2 and 3 ----
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
c = canvas.Canvas(P('Docs', 'Spec Sheet 480V UPS.pdf'), pagesize=letter)
c.setTitle('Spec Sheet 480V UPS')
pages = [
    ['480V Three-Phase UPS', 'Product overview', 'Rack and floor models from 100 to 500 kVA.'],
    ['Electrical', 'Input: 480V 3-phase, 60 Hz', 'Double-conversion efficiency up to 96.5 percent at half load.'],
    ['Operating modes', 'ECO mode efficiency reaches 99 percent.', 'Battery runtime: 8 minutes at full load.'],
]
for lines in pages:
    y = 720
    c.setFont('Helvetica-Bold', 20)
    c.drawString(72, y, lines[0])
    c.setFont('Helvetica', 12)
    for ln in lines[1:]:
        y -= 28
        c.drawString(72, y, ln)
    c.showPage()
c.save()

# ---- documents ----
from docx import Document
d = Document()
d.add_heading('Shop Rules', 1)
d.add_paragraph('Clean up after yourself. Put the torque wrench back at zero.')
d.save(P('Docs', 'Shop Rules.docx'))
from openpyxl import Workbook
wb = Workbook()
ws = wb.active
ws.title = 'Parts'
for r, row in enumerate([['Part', 'Price'], ['Oil filter', 9.99], ['Brake pads', 64.5], ['Wiper blades', 21]], 1):
    for col, v in enumerate(row, 1):
        ws.cell(r, col, v)
wb.save(P('Docs', 'Parts List.xlsx'))
write(P('Docs', 'readme.md'), '# The garage files\n\nEverything about the shop, in one folder.\n')
write(P('Docs', 'phone numbers.txt'), 'Parts store 555-0199\nTow truck 555-0123\n')


# ---- two WebM videos longer than five seconds, and two songs ----
def ff(*args):
    subprocess.run(['ffmpeg', '-v', 'error', '-y', *args], check=True)


ff('-f', 'lavfi', '-i', 'testsrc2=size=320x240:rate=24:duration=6', '-c:v', 'libvpx-vp9', '-deadline', 'realtime', '-cpu-used', '8', '-b:v', '150k', P('Videos', 'Garage Timelapse.webm'))
ff('-f', 'lavfi', '-i', 'testsrc=size=320x240:rate=24:duration=7', '-c:v', 'libvpx-vp9', '-deadline', 'realtime', '-cpu-used', '8', '-b:v', '150k', P('Videos', 'Beach Day.webm'))
ff('-f', 'lavfi', '-i', 'sine=frequency=330:sample_rate=44100:duration=4', '-metadata', 'artist=The Modem Tones', '-metadata', 'title=Handshake', '-metadata', 'album=56K', '-c:a', 'libmp3lame', '-b:a', '96k', P('Music', 'Handshake.mp3'))
ff('-f', 'lavfi', '-i', 'sine=frequency=550:sample_rate=44100:duration=3', '-metadata', 'artist=The Modem Tones', '-metadata', 'title=Carrier Lost', '-metadata', 'album=56K', '-c:a', 'libmp3lame', '-b:a', '96k', P('Music', 'Carrier Lost.mp3'))

n = sum(len(f) for _, _, f in os.walk(ROOT))
print('made', n, 'files in', ROOT)
